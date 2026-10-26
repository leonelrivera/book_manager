"""Capa de servicios y lógica de negocio para el sistema Book Manager.

Contiene los servicios encargados de implementar las reglas de negocio,
conversión de divisas, gestión de stock, orquestación de operaciones CRUD,
integración con APIs externas (DolarAPI) y generación de reportes analíticos.
"""

import csv
from datetime import date, datetime
import difflib
import functools
import os
import re
from typing import Any, Callable, Dict, List, Optional, Tuple, Union
import unicodedata
from dotenv import load_dotenv
import requests

from book_manager.entities.entities import (
    Auditoria,
    CotizacionDolar,
    Editorial,
    Genero,
    Libro,
    Moneda,
    Precio,
    Stock,
    TipoCotizacion,
)
from book_manager.models.models import AuditoriaModel
from book_manager.repositories.repositories import (
    RepositorioAuditoria,
    RepositorioBase,
    RepositorioCotizacionDolar,
    RepositorioEditorial,
    RepositorioGenero,
    RepositorioLibro,
    RepositorioMoneda,
    RepositorioPrecio,
    RepositorioStock,
    RepositorioTipoCotizacion,
)


_tabla_auditoria_verificada = False


def _persistir_auditoria(accion: str, fecha: datetime, detalles: str) -> None:
    """Inserta de manera segura una fila en la tabla de auditoría en la base de datos."""
    global _tabla_auditoria_verificada
    try:
        from book_manager.database.connection import ConexionDB, transaccion_db
        from book_manager.models.models import Base
        if not _tabla_auditoria_verificada:
            try:
                conexion = ConexionDB()
                Base.metadata.create_all(
                    bind=conexion.engine,
                    tables=[AuditoriaModel.__table__],
                    checkfirst=True,
                )
                _tabla_auditoria_verificada = True
            except Exception:
                pass

        with transaccion_db() as session:
            session.add(AuditoriaModel(
                accion=accion,
                fecha=fecha,
                detalles=detalles,
            ))
    except Exception:
        # Previene que un error transaccional de auditoría interrumpa la operación de negocio
        pass


def auditar(accion_o_func: Optional[Union[str, Callable]] = None) -> Any:
    """Decorador para registrar automáticamente auditorías de operaciones del sistema.

    Puede usarse con o sin argumentos explícitos:
        @auditar
        def crear_libro(self, ...): ...

        @auditar("ALTA_LIBRO")
        def crear_libro(self, ...): ...

    Registra en la tabla 'auditoria':
        - acción (accion): Nombre de la acción o nombre del método en mayúsculas
        - fecha (fecha): Timestamp de ejecución (datetime.now())
        - detalles (detalles): Parámetros de invocación, resultado obtenido y estado (ÉXITO / ERROR)
    """

    def decorador(func: Callable) -> Callable:
        @functools.wraps(func)
        def envoltura(*args: Any, **kwargs: Any) -> Any:
            if isinstance(accion_o_func, str):
                nombre_accion = accion_o_func
            else:
                nombre_accion = func.__name__.upper()

            fecha_actual = datetime.now()

            # Sanitización de parámetros excluyendo 'self'
            args_visibles = args[1:] if len(args) > 0 and hasattr(args[0], "__class__") else args
            partes = [repr(a)[:60] for a in args_visibles]
            partes += [f"{k}={repr(v)[:60]}" for k, v in kwargs.items()]
            params_str = ", ".join(partes) if partes else "ninguno"

            try:
                resultado = func(*args, **kwargs)
                detalles_retorno = ""
                if resultado is not None:
                    if hasattr(resultado, "id") and resultado.id:
                        detalles_retorno = f" | Retorno: ID={resultado.id}"
                    elif isinstance(resultado, (list, tuple)):
                        detalles_retorno = f" | Retorno: {len(resultado)} elementos"
                    elif isinstance(resultado, dict):
                        detalles_retorno = f" | Retorno: dict({len(resultado)} claves)"
                    else:
                        detalles_retorno = f" | Retorno: {str(resultado)[:60]}"

                detalles = f"Estado: ÉXITO | Parámetros: {params_str}{detalles_retorno}"
                _persistir_auditoria(nombre_accion, fecha_actual, detalles)
                return resultado
            except Exception as exc:
                detalles = f"Estado: ERROR ({type(exc).__name__}: {str(exc)}) | Parámetros: {params_str}"
                _persistir_auditoria(nombre_accion, fecha_actual, detalles)
                raise

        return envoltura

    if callable(accion_o_func):
        return decorador(accion_o_func)
    return decorador


class AuditoriaService:
    """Servicio de lógica de negocio para la consulta y administración de trazas de auditoría."""

    def __init__(self, repo_auditoria: RepositorioAuditoria) -> None:
        self._repo = repo_auditoria

    def registrar_auditoria(
        self,
        accion: str,
        detalles: str = "",
        fecha: Optional[datetime] = None,
    ) -> Auditoria:
        """Registra explícitamente un evento de auditoría."""
        nuevo = Auditoria(
            accion=accion,
            fecha=fecha or datetime.now(),
            detalles=detalles,
        )
        return self._repo.crear(nuevo)

    def listar_todas(self, limite: int = 100) -> List[Auditoria]:
        """Obtiene las últimas trazas de auditoría registradas."""
        return self._repo.leer_ultimas(limite=limite)

    def buscar_por_accion(self, accion: str) -> List[Auditoria]:
        """Filtra eventos de auditoría por tipo de acción."""
        return self._repo.buscar_por_accion(accion)

    def obtener_por_id(self, id_auditoria: int) -> Optional[Auditoria]:
        """Recupera un evento de auditoría por su identificador primario."""
        return self._repo.leer_por_id(id_auditoria)


class GeneroService:
    """Servicio de lógica de negocio para la gestión de Géneros literarios."""

    def __init__(self, repo_genero: RepositorioGenero) -> None:
        self._repo = repo_genero

    @auditar("CREAR_GENERO")
    def crear_genero(self, nombre: str, descripcion: str = "", id_genero: Optional[int] = None) -> Genero:
        if self._repo.buscar_por_nombre(nombre) is not None:
            raise ValueError(f"Ya existe un género literario con el nombre '{nombre}'.")
        nuevo_genero = Genero(id_genero=id_genero, nombre=nombre, descripcion=descripcion)
        return self._repo.crear(nuevo_genero)

    def obtener_por_id(self, id_genero: int) -> Optional[Genero]:
        return self._repo.leer_por_id(id_genero)

    def obtener_por_nombre(self, nombre: str) -> Optional[Genero]:
        return self._repo.buscar_por_nombre(nombre)

    def listar_todos(self) -> List[Genero]:
        return self._repo.leer_todos()

    @auditar("ACTUALIZAR_GENERO")
    def actualizar_genero(self, id_genero: int, nuevo_nombre: str, nueva_descripcion: str = "") -> Genero:
        genero_existente = self._repo.leer_por_id(id_genero)
        if not genero_existente:
            raise ValueError(f"No existe el género con ID {id_genero}.")
        otro_genero = self._repo.buscar_por_nombre(nuevo_nombre)
        if otro_genero and otro_genero.id != id_genero:
            raise ValueError(f"Ya existe otro género con el nombre '{nuevo_nombre}'.")
        genero_existente.nombre = nuevo_nombre
        genero_existente.descripcion = nueva_descripcion
        return self._repo.actualizar(genero_existente)

    @auditar("ELIMINAR_GENERO")
    def eliminar_genero(self, id_genero: int) -> bool:
        return self._repo.eliminar(id_genero)


class EditorialService:
    """Servicio de lógica de negocio para la gestión de Editoriales."""

    def __init__(self, repo_editorial: RepositorioEditorial) -> None:
        self._repo = repo_editorial

    @auditar("CREAR_EDITORIAL")
    def crear_editorial(
        self,
        nombre: str,
        pais: str = "Argentina",
        contacto: str = "",
        id_editorial: Optional[int] = None,
    ) -> Editorial:
        if self._repo.buscar_por_nombre(nombre) is not None:
            raise ValueError(f"Ya existe una editorial con el nombre '{nombre}'.")
        nueva_editorial = Editorial(id_editorial=id_editorial, nombre=nombre, pais=pais, contacto=contacto)
        return self._repo.crear(nueva_editorial)

    def obtener_por_id(self, id_editorial: int) -> Optional[Editorial]:
        return self._repo.leer_por_id(id_editorial)

    def obtener_por_nombre(self, nombre: str) -> Optional[Editorial]:
        return self._repo.buscar_por_nombre(nombre)

    def listar_todas(self) -> List[Editorial]:
        return self._repo.leer_todos()

    @auditar("ACTUALIZAR_EDITORIAL")
    def actualizar_editorial(
        self,
        id_editorial: int,
        nuevo_nombre: str,
        nuevo_pais: str = "Argentina",
        nuevo_contacto: str = "",
    ) -> Editorial:
        editorial_existente = self._repo.leer_por_id(id_editorial)
        if not editorial_existente:
            raise ValueError(f"No existe la editorial con ID {id_editorial}.")
        otra_editorial = self._repo.buscar_por_nombre(nuevo_nombre)
        if otra_editorial and otra_editorial.id != id_editorial:
            raise ValueError(f"Ya existe otra editorial con el nombre '{nuevo_nombre}'.")
        editorial_existente.nombre = nuevo_nombre
        editorial_existente.pais = nuevo_pais
        editorial_existente.contacto = nuevo_contacto
        return self._repo.actualizar(editorial_existente)

    @auditar("ELIMINAR_EDITORIAL")
    def eliminar_editorial(self, id_editorial: int) -> bool:
        return self._repo.eliminar(id_editorial)


class MonedaService:
    """Servicio de lógica de negocio para la gestión de Monedas."""

    def __init__(self, repo_moneda: RepositorioMoneda) -> None:
        self._repo = repo_moneda

    @auditar("CREAR_MONEDA")
    def crear_moneda(
        self,
        codigo: str,
        nombre: str,
        simbolo: str = "$",
        id_moneda: Optional[int] = None,
    ) -> Moneda:
        if self._repo.buscar_por_codigo(codigo) is not None:
            raise ValueError(f"Ya existe la moneda con código '{codigo}'.")
        nueva_moneda = Moneda(id_moneda=id_moneda, codigo=codigo, nombre=nombre, simbolo=simbolo)
        return self._repo.crear(nueva_moneda)

    def obtener_por_id(self, id_moneda: int) -> Optional[Moneda]:
        return self._repo.leer_por_id(id_moneda)

    def obtener_por_codigo(self, codigo: str) -> Optional[Moneda]:
        return self._repo.buscar_por_codigo(codigo)

    def listar_todas(self) -> List[Moneda]:
        return self._repo.leer_todos()

    @auditar("ACTUALIZAR_MONEDA")
    def actualizar_moneda(self, id_moneda: int, codigo: str, nombre: str, simbolo: str = "$") -> Moneda:
        moneda = self._repo.leer_por_id(id_moneda)
        if not moneda:
            raise ValueError(f"No existe la moneda con ID {id_moneda}.")
        otra = self._repo.buscar_por_codigo(codigo)
        if otra and otra.id != id_moneda:
            raise ValueError(f"Ya existe otra moneda con el código '{codigo}'.")
        moneda.codigo = codigo
        moneda.nombre = nombre
        moneda.simbolo = simbolo
        return self._repo.actualizar(moneda)

    @auditar("ELIMINAR_MONEDA")
    def eliminar_moneda(self, id_moneda: int) -> bool:
        return self._repo.eliminar(id_moneda)


class TipoCotizacionService:
    """Servicio de lógica de negocio para los Tipos de Cotización del dólar."""

    def __init__(self, repo_tipo: RepositorioTipoCotizacion) -> None:
        self._repo = repo_tipo

    @auditar("CREAR_TIPO_COTIZACION")
    def crear_tipo(self, nombre: str, descripcion: str = "", id_tipo: Optional[int] = None) -> TipoCotizacion:
        if self._repo.buscar_por_nombre(nombre) is not None:
            raise ValueError(f"Ya existe el tipo de cotización '{nombre}'.")
        nuevo_tipo = TipoCotizacion(id_tipo=id_tipo, nombre=nombre, descripcion=descripcion)
        return self._repo.crear(nuevo_tipo)

    def obtener_por_id(self, id_tipo: int) -> Optional[TipoCotizacion]:
        return self._repo.leer_por_id(id_tipo)

    def obtener_por_nombre(self, nombre: str) -> Optional[TipoCotizacion]:
        return self._repo.buscar_por_nombre(nombre)

    def listar_todos(self) -> List[TipoCotizacion]:
        return self._repo.leer_todos()

    @auditar("ACTUALIZAR_TIPO_COTIZACION")
    def actualizar_tipo(self, id_tipo: int, nuevo_nombre: str, nueva_descripcion: str = "") -> TipoCotizacion:
        tipo = self._repo.leer_por_id(id_tipo)
        if not tipo:
            raise ValueError(f"No existe el tipo de cotización con ID {id_tipo}.")
        otro = self._repo.buscar_por_nombre(nuevo_nombre)
        if otro and otro.id != id_tipo:
            raise ValueError(f"Ya existe otro tipo de cotización con el nombre '{nuevo_nombre}'.")
        tipo.nombre = nuevo_nombre
        tipo.descripcion = nueva_descripcion
        return self._repo.actualizar(tipo)

    @auditar("ELIMINAR_TIPO_COTIZACION")
    def eliminar_tipo(self, id_tipo: int) -> bool:
        return self._repo.eliminar(id_tipo)


class CotizacionService:
    """Servicio de lógica de negocio para el registro y consulta de cotizaciones del dólar."""

    def __init__(
        self,
        repo_cotizacion: RepositorioCotizacionDolar,
        repo_tipo: RepositorioTipoCotizacion,
    ) -> None:
        self._repo_cotizacion = repo_cotizacion
        self._repo_tipo = repo_tipo

    @auditar("REGISTRAR_COTIZACION")
    def registrar_cotizacion(
        self,
        id_tipo: int,
        fecha: Union[date, str],
        compra: float,
        venta: float,
        id_cotizacion: Optional[int] = None,
    ) -> CotizacionDolar:
        tipo = self._repo_tipo.leer_por_id(id_tipo)
        if not tipo:
            raise ValueError(f"El tipo de cotización con ID {id_tipo} no existe.")
        if compra < 0 or venta < 0:
            raise ValueError("Los valores de cotización no pueden ser negativos.")
        if compra > venta:
            raise ValueError("El valor de compra no puede ser mayor que el valor de venta.")

        nueva_cot = CotizacionDolar(
            id_cotizacion=id_cotizacion,
            tipo_cotizacion=tipo,
            fecha=fecha,
            compra=compra,
            venta=venta,
        )
        return self._repo_cotizacion.crear(nueva_cot)

    def obtener_ultima_cotizacion(self, id_tipo: int) -> Optional[CotizacionDolar]:
        return self._repo_cotizacion.leer_ultima_cotizacion(id_tipo)

    def obtener_historico_por_tipo(self, id_tipo: int) -> List[CotizacionDolar]:
        return self._repo_cotizacion.leer_historico_por_tipo(id_tipo)

    def listar_todas(self) -> List[CotizacionDolar]:
        return self._repo_cotizacion.leer_todos()

    @auditar("ELIMINAR_COTIZACION")
    def eliminar_cotizacion(self, id_tipo: int, fecha: Union[date, str]) -> bool:
        return self._repo_cotizacion.eliminar(id_tipo, fecha)

    @auditar("ACTUALIZAR_COTIZACIONES_API")
    def obtener_cotizaciones(self) -> List[CotizacionDolar]:
        """Busca, descarga y registra las cotizaciones actuales desde la API externa (DolarAPI).

        Lee la URL del endpoint exclusivamente desde la variable de entorno API_URL (.env)
        y actualiza la base de datos con las cotizaciones más recientes de cada casa de cambio.
        Distingue y gestiona errores específicos de red (Timeout, ConnectionError, HTTPError, RequestException).

        Returns:
            List[CotizacionDolar]: Lista de cotizaciones registradas/actualizadas.
        """
        load_dotenv()
        api_url = os.getenv("API_URL")
        if not api_url:
            print("❌ Error de configuración: La variable de entorno API_URL no está definida en el archivo .env.")
            return []

        try:
            respuesta = requests.get(api_url, timeout=10)
            respuesta.raise_for_status()
            datos = respuesta.json()
        except requests.exceptions.Timeout as e:
            print(f"⏳ Tiempo de espera agotado al conectar con la API de cotizaciones ({api_url}): {e}")
            return []
        except requests.exceptions.ConnectionError as e:
            print(f"🔌 Error de conexión de red al conectar con la API de cotizaciones ({api_url}): {e}")
            return []
        except requests.exceptions.HTTPError as e:
            print(f"🚫 Error de estado HTTP al consultar la API de cotizaciones ({api_url}): {e}")
            return []
        except requests.exceptions.RequestException as e:
            print(f"⚠️ Error en la solicitud hacia la API de cotizaciones ({api_url}): {e}")
            return []
        except ValueError as e:
            print(f"❌ Error al decodificar la respuesta JSON de la API ({api_url}): {e}")
            return []
        except Exception as e:
            print(f"❌ Error inesperado durante la obtención de cotizaciones ({api_url}): {e}")
            return []

        cotizaciones_registradas: List[CotizacionDolar] = []

        for item in datos:
            nombre_tipo = item.get("nombre", "").strip()
            if not nombre_tipo:
                continue

            compra = float(item.get("compra", 0.0) or 0.0)
            venta = float(item.get("venta", 0.0) or 0.0)

            # Normalizar fecha de actualización
            fecha_act_str = item.get("fechaActualizacion", "")
            if fecha_act_str:
                try:
                    fecha_obj = datetime.fromisoformat(fecha_act_str.replace("Z", "+00:00")).date()
                except Exception:
                    fecha_obj = date.today()
            else:
                fecha_obj = date.today()

            # Asegurar existencia del tipo de cotización
            tipo = self._repo_tipo.buscar_por_nombre(nombre_tipo)
            if tipo is None:
                tipo = TipoCotizacion(
                    id_tipo=None,
                    nombre=nombre_tipo,
                    descripcion=f"Cotización dólar {nombre_tipo} provista por DolarAPI",
                )
                tipo = self._repo_tipo.crear(tipo)

            # Registrar o actualizar la cotización en la base de datos
            cotizacion = CotizacionDolar(
                id_cotizacion=None,
                tipo_cotizacion=tipo,
                fecha=fecha_obj,
                compra=compra,
                venta=venta,
            )
            guardada = self._repo_cotizacion.crear(cotizacion)
            cotizaciones_registradas.append(guardada)

        return cotizaciones_registradas


# Alias solicitado por la consigna
ServicioCotizacionDolar = CotizacionService


class StockService:
    """Servicio de lógica de negocio para la administración de stock e inventario."""

    def __init__(self, repo_stock: RepositorioStock, repo_libro: RepositorioLibro) -> None:
        self._repo_stock = repo_stock
        self._repo_libro = repo_libro

    @auditar("REGISTRAR_STOCK_INICIAL")
    def registrar_stock_inicial(
        self,
        id_libro: int,
        cantidad: int = 0,
        ubicacion: str = "Depósito Central",
        id_stock: Optional[int] = None,
    ) -> Stock:
        libro = self._repo_libro.leer_por_id(id_libro)
        if not libro:
            raise ValueError(f"No existe el libro con ID {id_libro}.")
        if self._repo_stock.leer_por_libro(id_libro) is not None:
            raise ValueError(f"Ya existe un registro de stock para el libro '{libro.titulo}'.")
        nuevo_stock = Stock(id_stock=id_stock, libro=libro, cantidad=cantidad, ubicacion=ubicacion)
        return self._repo_stock.crear(nuevo_stock)

    def obtener_stock_por_libro(self, id_libro: int) -> Optional[Stock]:
        return self._repo_stock.leer_por_libro(id_libro)

    @auditar("AGREGAR_STOCK")
    def agregar_unidades(self, id_libro: int, cantidad: int) -> Stock:
        if cantidad <= 0:
            raise ValueError("La cantidad a ingresar debe ser mayor que cero.")
        stock = self._repo_stock.leer_por_libro(id_libro)
        if not stock:
            raise ValueError(f"No existe stock registrado para el libro ID {id_libro}.")
        stock.cantidad += cantidad
        return self._repo_stock.actualizar(stock)

    @auditar("RETIRAR_STOCK")
    def retirar_unidades(self, id_libro: int, cantidad: int) -> Stock:
        if cantidad <= 0:
            raise ValueError("La cantidad a retirar debe ser mayor que cero.")
        stock = self._repo_stock.leer_por_libro(id_libro)
        if not stock:
            raise ValueError(f"No existe stock registrado para el libro ID {id_libro}.")
        if stock.cantidad < cantidad:
            raise ValueError(
                f"Existencias insuficientes. Stock actual: {stock.cantidad}, solicitadas: {cantidad}."
            )
        stock.cantidad -= cantidad
        return self._repo_stock.actualizar(stock)

    def listar_todos(self) -> List[Stock]:
        return self._repo_stock.leer_todos()


class PrecioService:
    """Servicio de lógica de negocio para la cotización multimoneda y administración de precios."""

    def __init__(
        self,
        repo_precio: RepositorioPrecio,
        repo_libro: RepositorioLibro,
        repo_moneda: RepositorioMoneda,
        cotizacion_service: CotizacionService,
        tipo_cotizacion_service: TipoCotizacionService,
    ) -> None:
        self._repo_precio = repo_precio
        self._repo_libro = repo_libro
        self._repo_moneda = repo_moneda
        self._cotizacion_svc = cotizacion_service
        self._tipo_svc = tipo_cotizacion_service

    @auditar("FIJAR_PRECIO")
    def fijar_precio(
        self,
        id_libro: int,
        id_moneda: int,
        valor: float,
        id_precio: Optional[int] = None,
    ) -> Precio:
        if valor < 0:
            raise ValueError("El valor del precio no puede ser negativo.")
        libro = self._repo_libro.leer_por_id(id_libro)
        if not libro:
            raise ValueError(f"No existe el libro con ID {id_libro}.")
        moneda = self._repo_moneda.leer_por_id(id_moneda)
        if not moneda:
            raise ValueError(f"No existe la moneda con ID {id_moneda}.")

        existente = self._repo_precio.leer_por_libro_y_moneda(id_libro, id_moneda)
        if existente:
            existente.valor = valor
            return self._repo_precio.actualizar(existente)

        nuevo_precio = Precio(id_precio=id_precio, libro=libro, moneda=moneda, valor=valor)
        return self._repo_precio.crear(nuevo_precio)

    def obtener_precio_por_libro_y_moneda(self, id_libro: int, id_moneda: int) -> Optional[Precio]:
        return self._repo_precio.leer_por_libro_y_moneda(id_libro, id_moneda)

    def listar_precios_por_libro(self, id_libro: int) -> List[Precio]:
        return self._repo_precio.leer_por_libro(id_libro)

    def listar_todos(self) -> List[Precio]:
        return self._repo_precio.leer_todos()

    def convertir_precio(
        self,
        monto: float,
        moneda_origen: str,
        moneda_destino: str,
        tipo_dolar: str = "Blue",
    ) -> float:
        """Convierte montos monetarios considerando cotizaciones vigentes del dólar."""
        origen = moneda_origen.strip().upper()
        destino = moneda_destino.strip().upper()

        if origen == destino:
            return round(monto, 2)

        tipo = self._tipo_svc.obtener_por_nombre(tipo_dolar)
        if not tipo:
            raise ValueError(f"Tipo de cotización '{tipo_dolar}' no encontrado.")

        cotizacion = self._cotizacion_svc.obtener_ultima_cotizacion(tipo.id)
        if not cotizacion:
            raise ValueError(f"No existen cotizaciones registradas para el tipo '{tipo_dolar}'.")

        tasa_dolar = cotizacion.venta
        if tasa_dolar <= 0:
            raise ValueError("La tasa de cotización debe ser mayor que cero.")

        if origen == "ARS" and destino == "USD":
            return round(monto / tasa_dolar, 2)
        elif origen == "USD" and destino == "ARS":
            return round(monto * tasa_dolar, 2)
        else:
            raise NotImplementedError(f"Conversión entre {origen} y {destino} no soportada.")

    def obtener_precio_estimado(
        self,
        id_libro: int,
        moneda_destino: str,
        tipo_dolar: str = "Blue",
    ) -> Optional[float]:
        """Retorna el precio de un libro en la divisa solicitada realizando conversiones si es necesario."""
        moneda_dest_obj = self._repo_moneda.buscar_por_codigo(moneda_destino)
        if not moneda_dest_obj:
            return None

        precio_directo = self._repo_precio.leer_por_libro_y_moneda(id_libro, moneda_dest_obj.id)
        if precio_directo:
            return precio_directo.valor

        # Si no tiene precio directo, intentar convertir desde ARS o USD
        moneda_ars = self._repo_moneda.buscar_por_codigo("ARS")
        if moneda_ars and moneda_destino.upper() == "USD":
            precio_ars = self._repo_precio.leer_por_libro_y_moneda(id_libro, moneda_ars.id)
            if precio_ars:
                try:
                    return self.convertir_precio(precio_ars.valor, "ARS", "USD", tipo_dolar)
                except (ValueError, NotImplementedError):
                    return None

        moneda_usd = self._repo_moneda.buscar_por_codigo("USD")
        if moneda_usd and moneda_destino.upper() == "ARS":
            precio_usd = self._repo_precio.leer_por_libro_y_moneda(id_libro, moneda_usd.id)
            if precio_usd:
                try:
                    return self.convertir_precio(precio_usd.valor, "USD", "ARS", tipo_dolar)
                except (ValueError, NotImplementedError):
                    return None

        return None


class LibroService:
    """Servicio de lógica de negocio para la gestión integral del catálogo de Libros."""

    def __init__(
        self,
        repo_libro: RepositorioLibro,
        repo_genero: RepositorioGenero,
        repo_editorial: RepositorioEditorial,
    ) -> None:
        self._repo_libro = repo_libro
        self._repo_genero = repo_genero
        self._repo_editorial = repo_editorial

    @auditar("ALTA_LIBRO")
    def dar_de_alta_libro(
        self,
        isbn: str,
        titulo: str,
        autor: str,
        id_genero: int,
        id_editorial: int,
        anio_publicacion: Optional[int] = None,
        paginas: Optional[int] = None,
        id_libro: Optional[int] = None,
    ) -> Libro:
        if self._repo_libro.buscar_por_isbn(isbn) is not None:
            raise ValueError(f"Ya existe un libro registrado con el ISBN '{isbn}'.")

        genero = self._repo_genero.leer_por_id(id_genero)
        if not genero:
            raise ValueError(f"El género con ID {id_genero} no existe.")

        editorial = self._repo_editorial.leer_por_id(id_editorial)
        if not editorial:
            raise ValueError(f"La editorial con ID {id_editorial} no existe.")

        nuevo_libro = Libro(
            id_libro=id_libro,
            isbn=isbn,
            titulo=titulo,
            autor=autor,
            genero=genero,
            editorial=editorial,
            anio_publicacion=anio_publicacion,
            paginas=paginas,
        )
        return self._repo_libro.crear(nuevo_libro)

    def obtener_por_id(self, id_libro: int) -> Optional[Libro]:
        return self._repo_libro.leer_por_id(id_libro)

    def obtener_por_isbn(self, isbn: str) -> Optional[Libro]:
        return self._repo_libro.buscar_por_isbn(isbn)

    def buscar_por_titulo(self, termino: str) -> List[Libro]:
        return self._repo_libro.buscar_por_titulo(termino)

    def buscar_por_autor(self, autor: str) -> List[Libro]:
        return self._repo_libro.buscar_por_autor(autor)

    def buscar_por_genero(self, id_genero: int) -> List[Libro]:
        return self._repo_libro.buscar_por_genero(id_genero)

    def buscar_por_editorial(self, id_editorial: int) -> List[Libro]:
        return self._repo_libro.buscar_por_editorial(id_editorial)

    def listar_todos(self) -> List[Libro]:
        return self._repo_libro.leer_todos()

    @auditar("ACTUALIZAR_LIBRO")
    def actualizar_libro(
        self,
        id_libro: int,
        isbn: str,
        titulo: str,
        autor: str,
        id_genero: int,
        id_editorial: int,
        anio_publicacion: Optional[int] = None,
        paginas: Optional[int] = None,
    ) -> Libro:
        libro_existente = self._repo_libro.leer_por_id(id_libro)
        if not libro_existente:
            raise ValueError(f"No existe el libro con ID {id_libro}.")

        otro_isbn = self._repo_libro.buscar_por_isbn(isbn)
        if otro_isbn and otro_isbn.id != id_libro:
            raise ValueError(f"Ya existe otro libro registrado con el ISBN '{isbn}'.")

        genero = self._repo_genero.leer_por_id(id_genero)
        if not genero:
            raise ValueError(f"El género con ID {id_genero} no existe.")

        editorial = self._repo_editorial.leer_por_id(id_editorial)
        if not editorial:
            raise ValueError(f"La editorial con ID {id_editorial} no existe.")

        libro_existente.isbn = isbn
        libro_existente.titulo = titulo
        libro_existente.autor = autor
        libro_existente.genero = genero
        libro_existente.editorial = editorial
        libro_existente.anio_publicacion = anio_publicacion
        libro_existente.paginas = paginas

        return self._repo_libro.actualizar(libro_existente)

    @auditar("ELIMINAR_LIBRO")
    def eliminar_libro(self, id_libro: int) -> bool:
        return self._repo_libro.eliminar(id_libro)


class ReporteService:
    """Servicio de agregación y reportería analítica del catálogo e inventario."""

    def __init__(
        self,
        libro_service: LibroService,
        stock_service: StockService,
        precio_service: PrecioService,
        genero_service: GeneroService,
        editorial_service: EditorialService,
        cotizacion_service: CotizacionService,
        tipo_cotizacion_service: TipoCotizacionService,
        comparador_service: Optional["ComparadorPreciosService"] = None,
    ) -> None:
        self._libro_svc = libro_service
        self._stock_svc = stock_service
        self._precio_svc = precio_service
        self._genero_svc = genero_service
        self._editorial_svc = editorial_service
        self._cotizacion_svc = cotizacion_service
        self._tipo_svc = tipo_cotizacion_service
        self._comparador_service = comparador_service

    @auditar("REPORTE_EXCEL_PRECIOS")
    def generar_reporte_excel_precios(
        self,
        comparaciones: Optional[List[Dict[str, Any]]] = None,
        ruta_excel: Optional[str] = None,
    ) -> str:
        """Genera el reporte de precios en formato Excel (.xlsx) requerido en Ejercicio 05.

        Campos generados:
            - Libro: Título de la obra
            - Precio interno: Precio asignado en el catálogo interno en ARS
            - Precio web: Precio obtenido desde el competidor web (Cúspide Libros)
            - Diferencia: Discrepancia calculada (Precio web - Precio interno)
            - Fecha de extracción: Timestamp de captura del dato

        Args:
            comparaciones: Lista opcional de diccionarios de comparaciones.
            ruta_excel: Ruta opcional para guardar el archivo .xlsx.

        Returns:
            str: Ruta absoluta al archivo Excel generado.
        """
        if self._comparador_service is None:
            self._comparador_service = ComparadorPreciosService(self._libro_svc, self._precio_svc)
        return self._comparador_service.generar_reporte_excel(
            comparaciones=comparaciones,
            ruta_excel=ruta_excel,
        )

    def reporte_catalogo_completo(self, tipo_dolar: str = "Blue") -> List[Dict[str, Any]]:
        """Genera el catálogo consolidado de libros con existencias y precios en ARS y USD."""
        libros = self._libro_svc.listar_todos()
        resultado = []

        for libro in libros:
            genero = self._genero_svc.obtener_por_id(libro.id_genero)
            editorial = self._editorial_svc.obtener_por_id(libro.id_editorial)
            stock = self._stock_svc.obtener_stock_por_libro(libro.id)
            precio_ars = self._precio_svc.obtener_precio_estimado(libro.id, "ARS", tipo_dolar)
            precio_usd = self._precio_svc.obtener_precio_estimado(libro.id, "USD", tipo_dolar)

            resultado.append({
                "id": libro.id,
                "isbn": libro.isbn,
                "titulo": libro.titulo,
                "autor": libro.autor,
                "genero": genero.nombre if genero else "N/A",
                "editorial": editorial.nombre if editorial else "N/A",
                "stock": stock.cantidad if stock else 0,
                "precio_ars": precio_ars,
                "precio_usd": precio_usd,
            })
        return resultado

    def reporte_valorizacion_inventario(self, tipo_dolar: str = "Blue") -> Dict[str, Any]:
        """Calcula el valor total del inventario de libros en ARS y USD."""
        catalogo = self.reporte_catalogo_completo(tipo_dolar=tipo_dolar)
        total_unidades = sum(item["stock"] for item in catalogo)
        total_titulos = len(catalogo)
        valor_total_ars = sum((item["precio_ars"] or 0.0) * item["stock"] for item in catalogo)
        valor_total_usd = sum((item["precio_usd"] or 0.0) * item["stock"] for item in catalogo)

        return {
            "total_titulos": total_titulos,
            "total_ejemplares": total_unidades,
            "valor_total_ars": round(valor_total_ars, 2),
            "valor_total_usd": round(valor_total_usd, 2),
            "tipo_dolar_referencia": tipo_dolar,
        }

    def reporte_stock_critico(self, umbral: int = 5) -> List[Dict[str, Any]]:
        """Identifica libros con existencias inferiores o iguales al umbral establecido."""
        catalogo = self.reporte_catalogo_completo()
        return [item for item in catalogo if item["stock"] <= umbral]

    def reporte_distribucion_generos(self) -> Dict[str, int]:
        """Calcula la cantidad de libros disponibles por cada género."""
        distribucion: Dict[str, int] = {}
        for genero in self._genero_svc.listar_todos():
            distribucion[genero.nombre] = 0

        for libro in self._libro_svc.listar_todos():
            g = self._genero_svc.obtener_por_id(libro.id_genero)
            nombre_genero = g.nombre if g else "Sin Género"
            distribucion[nombre_genero] = distribucion.get(nombre_genero, 0) + 1

        return distribucion


class ComparadorPreciosService:
    """Servicio de lógica de negocio para la comparación de precios internos vs competencia (Cúspide).

    Implementa el contraste entre precios internos del catálogo y los precios
    extraídos mediante web scraping, evaluando umbrales de diferencia ingresados
    por el usuario y generando reportes de alertas en formato CSV descargable.
    """

    def __init__(self, libro_service: LibroService, precio_service: PrecioService) -> None:
        """Inicializa el comparador de precios vinculando los servicios de libros y precios.

        Args:
            libro_service: Servicio de libros para consulta de catálogo.
            precio_service: Servicio de precios para consulta de valores internos en ARS.
        """
        self._libro_service = libro_service
        self._precio_service = precio_service

    @auditar("COMPARAR_PRECIOS")
    def comparar_precios(
        self,
        items_scraped: List[Dict[str, Any]],
        umbral_diferencia: float = 0.0,
    ) -> List[Dict[str, Any]]:
        """Compara los precios extraídos de la web contra el catálogo interno.

        Args:
            items_scraped: Lista de diccionarios de ítems obtenidos por CuspideSpider.
            umbral_diferencia: Valor mínimo en ARS de discrepancia de precio para
                generar una alerta (ingresado por el usuario). Por defecto 0.0.

        Returns:
            List[Dict[str, Any]]: Lista de registros de comparación con estado de alerta.
        """
        comparaciones: List[Dict[str, Any]] = []
        umbral = float(abs(umbral_diferencia))

        libros_propios = self._libro_service.listar_todos()

        def _norm(txt: Optional[str]) -> str:
            if not txt:
                return ""
            nfkd = unicodedata.normalize("NFKD", str(txt))
            sin_tildes = "".join(c for c in nfkd if not unicodedata.combining(c))
            return re.sub(r"[^a-z0-9\s]", " ", sin_tildes.lower()).strip()

        # Agrupar ítems extraídos validados por cada libro de nuestro catálogo
        candidatos_por_libro: Dict[int, List[Tuple[float, Dict[str, Any]]]] = {
            lib.id: [] for lib in libros_propios if lib.id is not None
        }

        for item in items_scraped:
            titulo_int = str(item.get("titulo_interno", "")).strip()
            titulo_comp = str(item.get("titulo_competidor", "")).strip()
            t_int_norm = _norm(titulo_int)
            t_comp_norm = _norm(titulo_comp)

            # Encontrar el libro correspondiente en el catálogo
            libro_match = None
            for lib in libros_propios:
                if lib.titulo and _norm(lib.titulo) == t_int_norm:
                    libro_match = lib
                    break
            if not libro_match:
                for lib in libros_propios:
                    if lib.titulo and (_norm(lib.titulo) in t_int_norm or t_int_norm in _norm(lib.titulo)):
                        libro_match = lib
                        break

            if libro_match and libro_match.id is not None:
                # Calcular similitud textual para elegir la edición principal
                ratio = difflib.SequenceMatcher(None, _norm(libro_match.titulo), t_comp_norm).ratio()
                candidatos_por_libro[libro_match.id].append((ratio, item))

        for lib in libros_propios:
            if lib.id is None:
                continue
            isbn = lib.isbn if lib.isbn else "N/D"
            id_libro = lib.id
            titulo_int = lib.titulo

            # Obtener precio interno en Pesos Argentinos (id_moneda=1)
            precio_interno: Optional[float] = None
            p_obj = self._precio_service.obtener_precio_por_libro_y_moneda(id_libro, id_moneda=1)
            if p_obj and p_obj.valor is not None:
                precio_interno = float(p_obj.valor)

            candidatos = candidatos_por_libro.get(id_libro, [])
            if candidatos:
                # Ordenar por mayor ratio de similitud para elegir el mejor match
                candidatos.sort(key=lambda x: x[0], reverse=True)
                _, mejor_item = candidatos[0]

                titulo_comp = str(mejor_item.get("titulo_competidor", "")).strip()
                precio_web = float(mejor_item["precio"]) if mejor_item.get("precio") is not None else None
                url_img = mejor_item.get("url_imagen", "")
                formas_pago = mejor_item.get("formas_pago", "")
                url_libro = mejor_item.get("url_libro", "")
                fecha_ext = mejor_item.get("fecha_extraccion", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            else:
                titulo_comp = "Sin coincidencia en Cúspide"
                precio_web = None
                url_img = ""
                formas_pago = ""
                url_libro = ""
                fecha_ext = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            # Cálculo de diferencias y determinación del tipo de alerta
            if precio_web is not None and precio_interno is not None:
                diferencia = round(precio_web - precio_interno, 2)
                diferencia_abs = abs(diferencia)
                porcentaje = (
                    round((diferencia / precio_interno) * 100, 2) if precio_interno > 0 else 0.0
                )

                if diferencia < 0:
                    tipo_alerta = "COMPETIDOR MÁS BARATO"
                    detalle = f"Cúspide vende ${diferencia_abs:,.2f} ({abs(porcentaje)}%) por debajo del precio interno"
                elif diferencia > 0:
                    tipo_alerta = "COMPETIDOR MÁS CARO"
                    detalle = f"Cúspide vende ${diferencia:,.2f} (+{porcentaje}%) por encima del precio interno"
                else:
                    tipo_alerta = "PRECIOS IGUALES"
                    detalle = "El precio de Cúspide coincide exactamente con el precio interno"

                alerta_activa = diferencia_abs >= umbral
            else:
                diferencia = 0.0
                diferencia_abs = 0.0
                porcentaje = 0.0
                if precio_interno is None and precio_web is not None:
                    tipo_alerta = "SIN PRECIO INTERNO"
                    detalle = "El libro no tiene precio asignado en el catálogo interno"
                elif precio_web is None and precio_interno is not None:
                    tipo_alerta = "SIN PRECIO WEB"
                    detalle = "No se encontró edición coincidente en Cúspide Libros"
                else:
                    tipo_alerta = "SIN PRECIOS"
                    detalle = "Sin precio interno ni web disponible"
                alerta_activa = False

            registro: Dict[str, Any] = {
                "id_libro": id_libro,
                "isbn": isbn,
                "titulo_interno": titulo_int,
                "titulo_competidor": titulo_comp,
                "precio_interno": precio_interno,
                "precio_web": precio_web,
                "diferencia": diferencia,
                "diferencia_abs": diferencia_abs,
                "diferencia_porcentual": porcentaje,
                "tipo_alerta": tipo_alerta,
                "detalle_alerta": detalle,
                "alerta_activa": alerta_activa,
                "formas_pago": formas_pago,
                "url_imagen": url_img,
                "url_libro": url_libro,
                "fecha_extraccion": fecha_ext,
            }
            comparaciones.append(registro)

        return comparaciones

    @auditar("GENERAR_ALERTAS_CSV")
    def generar_alertas_csv(
        self,
        alertas: List[Dict[str, Any]],
        ruta_csv: Optional[str] = None,
        solo_alertas_activas: bool = False,
    ) -> str:
        """Genera un archivo CSV con las alertas de precios para su descarga y análisis.

        Args:
            alertas: Lista de diccionarios generados por comparar_precios.
            ruta_csv: Ruta opcional de destino. Por defecto se guarda en 'reportes/alertas_precios.csv'.
            solo_alertas_activas: Si es True, filtra únicamente registros con alerta_activa=True.

        Returns:
            str: Ruta absoluta al archivo CSV generado.
        """
        if ruta_csv is None:
            directorio_base = os.path.abspath(
                os.path.join(os.path.dirname(__file__), "..", "reportes")
            )
            os.makedirs(directorio_base, exist_ok=True)
            ruta_csv = os.path.join(directorio_base, "alertas_precios.csv")
        else:
            directorio_destino = os.path.dirname(os.path.abspath(ruta_csv))
            if directorio_destino:
                os.makedirs(directorio_destino, exist_ok=True)

        campos = [
            "isbn",
            "titulo_interno",
            "titulo_competidor",
            "precio_interno",
            "precio_web",
            "diferencia",
            "diferencia_porcentual",
            "alerta_activa",
            "tipo_alerta",
            "detalle_alerta",
            "formas_pago",
            "url_libro",
            "fecha_extraccion",
        ]

        filas_a_escribir = [
            a for a in alertas if (not solo_alertas_activas or a.get("alerta_activa", False))
        ] if alertas else []

        with open(ruta_csv, mode="w", encoding="utf-8", newline="") as f:
            escritor = csv.DictWriter(f, fieldnames=campos, extrasaction="ignore")
            escritor.writeheader()
            for fila in filas_a_escribir:
                fila_csv = dict(fila)
                fila_csv["alerta_activa"] = "SÍ" if fila.get("alerta_activa") else "NO"
                p_int = fila_csv.get("precio_interno")
                fila_csv["precio_interno"] = f"{p_int:.2f}" if p_int is not None else "N/D"
                p_web = fila_csv.get("precio_web")
                fila_csv["precio_web"] = f"{p_web:.2f}" if p_web is not None else "N/D"
                dif = fila_csv.get("diferencia")
                fila_csv["diferencia"] = f"{dif:.2f}" if dif is not None else "0.00"
                dif_pct = fila_csv.get("diferencia_porcentual")
                fila_csv["diferencia_porcentual"] = f"{dif_pct:.2f}%" if dif_pct is not None else "0.00%"
                escritor.writerow(fila_csv)

        return os.path.abspath(ruta_csv)

    @auditar("SCRAPING_Y_ALERTAS")
    def ejecutar_scraping_y_alertas(
        self,
        umbral_diferencia: float = 0.0,
        libros: Optional[List[str]] = None,
        ruta_csv: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Ejecuta el scraper de Cúspide y genera automáticamente el CSV de alertas.

        Args:
            umbral_diferencia: Diferencia mínima requerida para generar alertas.
            libros: Lista opcional de títulos de libros a scrapear.
            ruta_csv: Ruta opcional para guardar el archivo CSV de alertas.

        Returns:
            Dict[str, Any]: Resumen con total extraído, total alertas y ruta del CSV.
        """
        from book_manager.scraper.runner import ejecutar_scraper

        print(f"🕷️ Iniciando scraping de competidor (Cúspide) con umbral de diferencia: ${umbral_diferencia:,.2f}...")
        items_scraped = ejecutar_scraper(libros=libros)
        print(f"📦 Scraping finalizado. {len(items_scraped)} registros obtenidos de Cúspide.")

        alertas = self.comparar_precios(items_scraped, umbral_diferencia=umbral_diferencia)
        ruta_generada = self.generar_alertas_csv(alertas, ruta_csv=ruta_csv)

        total_activas = sum(1 for a in alertas if a.get("alerta_activa", False))
        print(f"⚠️ Comparación completada: {total_activas} alertas activas detectadas.")
        print(f"📄 Archivo CSV de alertas disponible en: {ruta_generada}")

        return {
            "items_scraped": items_scraped,
            "alertas": alertas,
            "total_extraidos": len(items_scraped),
            "total_alertas": total_activas,
            "archivo_csv": ruta_generada,
        }

    @auditar("GENERAR_REPORTE_EXCEL")
    def generar_reporte_excel(
        self,
        comparaciones: Optional[List[Dict[str, Any]]] = None,
        ruta_excel: Optional[str] = None,
    ) -> str:
        """Genera un reporte analítico en formato Excel (.xlsx) con openpyxl.

        Campos requeridos según especificación de Ejercicio 05:
            - Libro: Título de la obra literaria
            - Precio interno: Precio asignado en el catálogo interno en ARS
            - Precio web: Precio extraído del competidor web (Cúspide Libros)
            - Diferencia: Discrepancia calculada (Precio web - Precio interno)
            - Fecha de extracción: Timestamp de captura del dato

        Args:
            comparaciones: Lista opcional de diccionarios con comparaciones.
                Si no se proporciona, intenta leer desde reportes/alertas_precios.csv
                o genera comparaciones a partir del catálogo interno.
            ruta_excel: Ruta de destino opcional del archivo Excel. Por defecto
                se guarda en 'reportes/reporte_precios.xlsx'.

        Returns:
            str: Ruta absoluta al archivo Excel generado.
        """
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        if ruta_excel is None:
            directorio_base = os.path.abspath(
                os.path.join(os.path.dirname(__file__), "..", "reportes")
            )
            os.makedirs(directorio_base, exist_ok=True)
            ruta_excel = os.path.join(directorio_base, "reporte_precios.xlsx")
        else:
            directorio_destino = os.path.dirname(os.path.abspath(ruta_excel))
            if directorio_destino:
                os.makedirs(directorio_destino, exist_ok=True)

        # Si no se suministraron comparaciones, cargar desde alertas_precios.csv o catálogo
        if not comparaciones:
            directorio_base = os.path.dirname(os.path.abspath(ruta_excel))
            ruta_csv_alertas = os.path.join(directorio_base, "alertas_precios.csv")
            if os.path.exists(ruta_csv_alertas):
                comparaciones = []
                with open(ruta_csv_alertas, mode="r", encoding="utf-8") as f:
                    lector = csv.DictReader(f)
                    for fila in lector:
                        p_int_raw = fila.get("precio_interno")
                        p_web_raw = fila.get("precio_web")
                        dif_raw = fila.get("diferencia")
                        try:
                            p_int = float(p_int_raw) if p_int_raw not in (None, "", "N/D") else None
                        except ValueError:
                            p_int = None
                        try:
                            p_web = float(p_web_raw) if p_web_raw not in (None, "", "N/D") else None
                        except ValueError:
                            p_web = None
                        try:
                            dif = float(dif_raw) if dif_raw not in (None, "", "N/D") else 0.0
                        except ValueError:
                            dif = 0.0
                        comparaciones.append({
                            "titulo_interno": fila.get("titulo_interno") or fila.get("titulo_competidor") or "Desconocido",
                            "precio_interno": p_int,
                            "precio_web": p_web,
                            "diferencia": dif,
                            "fecha_extraccion": fila.get("fecha_extraccion") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        })
            else:
                libros = self._libro_service.listar_todos()
                items_mock = []
                for lib in libros:
                    items_mock.append({
                        "titulo_interno": lib.titulo,
                        "titulo_competidor": f"{lib.titulo.upper()} - {lib.autor.upper()}",
                        "precio": None,
                        "fecha_extraccion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                comparaciones = self.comparar_precios(items_mock)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte de Precios"
        ws.views.sheetView[0].showGridLines = True

        # Encabezado del reporte
        ws.merge_cells("A1:E1")
        celda_titulo = ws["A1"]
        celda_titulo.value = "BOOK MANAGER - REPORTE DE PRECIOS Y COMPETENCIA (CÚSPIDE)"
        celda_titulo.font = Font(name="Calibri", size=13, bold=True, color="1F497D")
        celda_titulo.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

        ws.merge_cells("A2:E2")
        celda_subtitulo = ws["A2"]
        celda_subtitulo.value = (
            f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | "
            f"Total registros: {len(comparaciones)} | Fuente: Cúspide Libros Online"
        )
        celda_subtitulo.font = Font(name="Calibri", size=9, italic=True, color="595959")
        celda_subtitulo.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 8

        # Fila 4: Columnas oficiales del Ejercicio 05
        columnas = [
            "Libro",
            "Precio interno",
            "Precio web",
            "Diferencia",
            "Fecha de extracción",
        ]
        ws.row_dimensions[4].height = 24

        fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        relleno_encabezado = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        alineacion_encabezado = Alignment(horizontal="center", vertical="center", wrap_text=True)

        borde_fino = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_idx, nom_col in enumerate(columnas, start=1):
            c = ws.cell(row=4, column=col_idx, value=nom_col)
            c.font = fuente_encabezado
            c.fill = relleno_encabezado
            c.alignment = alineacion_encabezado
            c.border = borde_fino

        fuente_dato = Font(name="Calibri", size=10)
        relleno_par = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        relleno_impar = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        formato_moneda = "$#,##0.00"
        formato_dif = "$#,##0.00;[Red]($#,##0.00);\"$0.00\""

        fila_inicio = 5
        for i, comp in enumerate(comparaciones, start=fila_inicio):
            titulo_libro = comp.get("titulo_interno") or comp.get("titulo_competidor") or "Desconocido"
            p_int = comp.get("precio_interno")
            p_web = comp.get("precio_web")
            dif = comp.get("diferencia")
            fecha_ext = str(comp.get("fecha_extraccion") or datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

            c_libro = ws.cell(row=i, column=1, value=titulo_libro)
            c_libro.alignment = Alignment(horizontal="left", vertical="center")

            c_pint = ws.cell(row=i, column=2, value=p_int if p_int is not None else "N/D")
            c_pint.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(p_int, (int, float)):
                c_pint.number_format = formato_moneda

            c_pweb = ws.cell(row=i, column=3, value=p_web if p_web is not None else "N/D")
            c_pweb.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(p_web, (int, float)):
                c_pweb.number_format = formato_moneda

            c_dif = ws.cell(row=i, column=4, value=dif if dif is not None else 0.0)
            c_dif.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(dif, (int, float)):
                c_dif.number_format = formato_dif

            c_fecha = ws.cell(row=i, column=5, value=fecha_ext)
            c_fecha.alignment = Alignment(horizontal="center", vertical="center")

            relleno_actual = relleno_par if (i % 2 == 0) else relleno_impar
            for col_idx in range(1, 6):
                celda = ws.cell(row=i, column=col_idx)
                celda.font = fuente_dato
                celda.fill = relleno_actual
                celda.border = borde_fino

            ws.row_dimensions[i].height = 20

        fila_fin = fila_inicio + len(comparaciones) - 1
        fila_total = fila_fin + 1
        if len(comparaciones) > 0:
            c_tot_lbl = ws.cell(row=fila_total, column=1, value="PROMEDIO GENERAL")
            c_tot_lbl.font = Font(name="Calibri", size=10, bold=True, color="1F497D")
            c_tot_lbl.alignment = Alignment(horizontal="left", vertical="center")

            c_tot_pint = ws.cell(row=fila_total, column=2, value=f"=AVERAGE(B{fila_inicio}:B{fila_fin})")
            c_tot_pint.font = Font(name="Calibri", size=10, bold=True, color="1F497D")
            c_tot_pint.number_format = formato_moneda
            c_tot_pint.alignment = Alignment(horizontal="right", vertical="center")

            c_tot_pweb = ws.cell(row=fila_total, column=3, value=f"=AVERAGE(C{fila_inicio}:C{fila_fin})")
            c_tot_pweb.font = Font(name="Calibri", size=10, bold=True, color="1F497D")
            c_tot_pweb.number_format = formato_moneda
            c_tot_pweb.alignment = Alignment(horizontal="right", vertical="center")

            c_tot_dif = ws.cell(row=fila_total, column=4, value=f"=AVERAGE(D{fila_inicio}:D{fila_fin})")
            c_tot_dif.font = Font(name="Calibri", size=10, bold=True, color="1F497D")
            c_tot_dif.number_format = formato_dif
            c_tot_dif.alignment = Alignment(horizontal="right", vertical="center")

            c_tot_fec = ws.cell(row=fila_total, column=5, value=f"Total: {len(comparaciones)} títulos")
            c_tot_fec.font = Font(name="Calibri", size=9, bold=True, italic=True, color="595959")
            c_tot_fec.alignment = Alignment(horizontal="center", vertical="center")

            borde_total = Border(
                top=Side(style="thin", color="1F497D"),
                bottom=Side(style="double", color="1F497D"),
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
            )
            relleno_total = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")

            for col_idx in range(1, 6):
                celda = ws.cell(row=fila_total, column=col_idx)
                celda.fill = relleno_total
                celda.border = borde_total

            ws.row_dimensions[fila_total].height = 22

        anchos_minimos = {1: 32, 2: 16, 3: 16, 4: 16, 5: 22}
        for col_idx in range(1, 6):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(4, ws.max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    val_str = str(val)
                    if not val_str.startswith("="):
                        max_len = max(max_len, len(val_str))
            ancho_final = max(max_len + 3, anchos_minimos.get(col_idx, 15))
            ws.column_dimensions[col_letter].width = ancho_final

        wb.save(ruta_excel)
        return os.path.abspath(ruta_excel)


class BookManagerService:
    """Coordinador central que implementa inyección de dependencias modular.

    Permite instanciar repositorios y servicios externamente o provee
    instancias predeterminadas conectadas a la base de datos relacional.
    """

    def __init__(
        self,
        repo_libro: Optional[RepositorioLibro] = None,
        repo_genero: Optional[RepositorioGenero] = None,
        repo_editorial: Optional[RepositorioEditorial] = None,
        repo_moneda: Optional[RepositorioMoneda] = None,
        repo_tipo_cotizacion: Optional[RepositorioTipoCotizacion] = None,
        repo_precio: Optional[RepositorioPrecio] = None,
        repo_stock: Optional[RepositorioStock] = None,
        repo_cotizacion: Optional[RepositorioCotizacionDolar] = None,
        repo_auditoria: Optional[RepositorioAuditoria] = None,
    ) -> None:
        # Inicialización de Repositorios con Inyección de Dependencias
        self.repo_libro = repo_libro or RepositorioLibro()
        self.repo_genero = repo_genero or RepositorioGenero()
        self.repo_editorial = repo_editorial or RepositorioEditorial()
        self.repo_moneda = repo_moneda or RepositorioMoneda()
        self.repo_tipo_cotizacion = repo_tipo_cotizacion or RepositorioTipoCotizacion()
        self.repo_precio = repo_precio or RepositorioPrecio()
        self.repo_stock = repo_stock or RepositorioStock()
        self.repo_cotizacion = repo_cotizacion or RepositorioCotizacionDolar()
        self.repo_auditoria = repo_auditoria or RepositorioAuditoria()

        # Inicialización de Servicios
        self.auditoria_service = AuditoriaService(self.repo_auditoria)
        self.genero_service = GeneroService(self.repo_genero)
        self.editorial_service = EditorialService(self.repo_editorial)
        self.moneda_service = MonedaService(self.repo_moneda)
        self.tipo_cotizacion_service = TipoCotizacionService(self.repo_tipo_cotizacion)
        self.cotizacion_service = CotizacionService(self.repo_cotizacion, self.repo_tipo_cotizacion)
        self.libro_service = LibroService(self.repo_libro, self.repo_genero, self.repo_editorial)
        self.stock_service = StockService(self.repo_stock, self.repo_libro)
        self.precio_service = PrecioService(
            self.repo_precio,
            self.repo_libro,
            self.repo_moneda,
            self.cotizacion_service,
            self.tipo_cotizacion_service,
        )
        self.comparador_service = ComparadorPreciosService(
            self.libro_service,
            self.precio_service,
        )
        self.reporte_service = ReporteService(
            self.libro_service,
            self.stock_service,
            self.precio_service,
            self.genero_service,
            self.editorial_service,
            self.cotizacion_service,
            self.tipo_cotizacion_service,
            self.comparador_service,
        )
